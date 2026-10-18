"""
SkillGraph AI — Privacy-Preserving Résumé Skill Extraction Service.

Extracts candidate skills from résumé text or uploaded files using NLP pattern matching
against the canonical skills taxonomy (8,828 skills).
Includes explicit privacy compliance: processing is strictly in-memory and ephemeral.
"""

from __future__ import annotations

import logging
import re
import time
from typing import Any, Dict, List, Optional, Set
from sqlalchemy.orm import Session

from backend.app.db.models import Skill

logger = logging.getLogger(__name__)


class ResumeParserService:
    """Extracts candidate skills from résumé text with zero permanent storage."""

    def __init__(self):
        self._skills_cache: Optional[List[Dict[str, Any]]] = None

    def _get_skills_dictionary(self, db: Session) -> List[Dict[str, Any]]:
        """Load and cache canonical skills with regex patterns."""
        if self._skills_cache is not None:
            return self._skills_cache

        skills = db.query(Skill.id, Skill.canonical_name, Skill.category).all()
        cache = []
        for s_id, name, cat in skills:
            if not name or len(name) < 2:
                continue
            # Compile word-boundary regex pattern
            pattern = re.compile(rf"\b{re.escape(name)}\b", re.IGNORECASE)
            cache.append({
                "id": s_id,
                "name": name,
                "category": cat or "Technical",
                "pattern": pattern,
            })

        self._skills_cache = cache
        logger.info(f"Loaded {len(cache)} skill patterns for résumé parsing.")
        return self._skills_cache

    def extract_from_text(
        self,
        db: Session,
        text: str,
        user_experience_hint: Optional[str] = None,
    ) -> Dict[str, Any]:
        """
        Extract candidate competencies from raw résumé text.
        Guarantees zero data retention.
        """
        start_time = time.time()

        if not text or len(text.strip()) < 10:
            return {
                "extracted_skills": [],
                "total_skills_found": 0,
                "word_count": len(text.split()) if text else 0,
                "processing_time_ms": 0.0,
                "message": "Text too short to extract meaningful skills.",
                "privacy_notice": "In-memory processing only. No data was retained.",
            }

        skills_dict = self._get_skills_dictionary(db)
        word_count = len(text.split())

        detected: List[Dict[str, Any]] = []
        seen_names: Set[str] = set()

        for s in skills_dict:
            matches = list(s["pattern"].finditer(text))
            if matches:
                name = s["name"]
                if name.lower() in seen_names:
                    continue
                seen_names.add(name.lower())

                # Find snippet context from first occurrence
                m = matches[0]
                start = max(0, m.start() - 35)
                end = min(len(text), m.end() + 35)
                snippet = "..." + text[start:end].replace("\n", " ").strip() + "..."

                # Infer candidate proficiency based on occurrences & context
                occ_count = len(matches)
                if occ_count >= 3 or any(w in snippet.lower() for w in ["lead", "architect", "expert", "senior", "advanced"]):
                    proficiency = "advanced"
                    confidence = 0.95
                elif occ_count >= 2 or any(w in snippet.lower() for w in ["developed", "built", "implemented", "engineered"]):
                    proficiency = "intermediate"
                    confidence = 0.88
                else:
                    proficiency = "beginner"
                    confidence = 0.78

                detected.append({
                    "skill_id": s["id"],
                    "skill_name": name,
                    "category": s["category"],
                    "proficiency": proficiency,
                    "confidence": confidence,
                    "occurrences": occ_count,
                    "context_snippet": snippet,
                })

        # Sort by confidence and occurrences descending
        detected.sort(key=lambda x: (x["confidence"], x["occurrences"]), reverse=True)

        # Experience level detection
        detected_exp = user_experience_hint or "student"
        lower_full = text.lower()
        if any(w in lower_full for w in ["5+ years", "6+ years", "lead architect", "principal engineer", "director"]):
            detected_exp = "senior"
        elif any(w in lower_full for w in ["2+ years", "3+ years", "mid-level", "software engineer ii"]):
            detected_exp = "mid_level"
        elif any(w in lower_full for w in ["student", "intern", "fresh graduate", "bachelor", "entry level"]):
            detected_exp = "student"

        elapsed_ms = round((time.time() - start_time) * 1000, 1)

        return {
            "extracted_skills": detected,
            "total_skills_found": len(detected),
            "word_count": word_count,
            "inferred_experience_level": detected_exp,
            "processing_time_ms": elapsed_ms,
            "privacy_notice": (
                "Privacy Guarantee: Your résumé text was parsed entirely in-memory "
                "for skill detection and has been discarded. No documents or personal identifiers "
                "are permanently retained or shared."
            ),
        }

    def extract_from_file_bytes(
        self,
        db: Session,
        file_bytes: bytes,
        filename: str,
        user_experience_hint: Optional[str] = None,
    ) -> Dict[str, Any]:
        """
        Extract text from PDF, DOCX, DOC, or TXT file bytes and parse competencies.
        Guarantees zero persistent storage of the uploaded file.
        """
        lower_name = filename.lower()
        extracted_text = ""

        # 1. PDF extraction
        if lower_name.endswith(".pdf"):
            try:
                import io
                import pypdf
                reader = pypdf.PdfReader(io.BytesIO(file_bytes))
                pages_text = []
                for p in reader.pages:
                    txt = p.extract_text()
                    if txt:
                        pages_text.append(txt)
                extracted_text = "\n".join(pages_text)
            except Exception as e:
                logger.warning(f"Error parsing PDF with pypdf: {e}")
                extracted_text = re.sub(r"[^\x20-\x7E\n]", " ", file_bytes.decode("latin-1", errors="ignore"))

        # 2. DOCX extraction
        elif lower_name.endswith(".docx"):
            try:
                import io
                import docx
                doc = docx.Document(io.BytesIO(file_bytes))
                paragraphs = [p.text for p in doc.paragraphs if p.text]
                for table in doc.tables:
                    for row in table.rows:
                        paragraphs.extend(cell.text for cell in row.cells if cell.text)
                extracted_text = "\n".join(paragraphs)
            except Exception as e:
                logger.warning(f"Error parsing DOCX with python-docx: {e}")
                extracted_text = re.sub(r"[^\x20-\x7E\n]", " ", file_bytes.decode("latin-1", errors="ignore"))

        # 3. DOC (legacy binary Word) or fallback
        elif lower_name.endswith(".doc"):
            try:
                text_stream = file_bytes.decode("utf-8", errors="ignore")
                extracted_text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", " ", text_stream)
            except Exception:
                extracted_text = re.sub(r"[^\x20-\x7E\n]", " ", file_bytes.decode("latin-1", errors="ignore"))

        # 4. Plain text / Markdown
        else:
            try:
                extracted_text = file_bytes.decode("utf-8")
            except UnicodeDecodeError:
                extracted_text = file_bytes.decode("latin-1", errors="ignore")

        res = self.extract_from_text(db, text=extracted_text, user_experience_hint=user_experience_hint)
        res["source_filename"] = filename
        return res


resume_service = ResumeParserService()
